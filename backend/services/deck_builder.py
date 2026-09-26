"""Build a branded sales deck from analysis JSON. Product claims stay blank."""

from __future__ import annotations

import logging
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

logger = logging.getLogger(__name__)

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "assets" / "template.pptx"
SLIDE_W = Inches(13.333333)
SLIDE_H = Inches(7.5)
FONT = "Inter"
INK = RGBColor(0x1C, 0x24, 0x33)
MUTED = RGBColor(0x5C, 0x67, 0x75)
PAPER = RGBColor(0xF7, 0xF5, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xE7, 0xEE, 0xF2)


def _rgb(value: str) -> RGBColor:
    color = value.strip().lstrip("#")
    if len(color) == 3:
        color = "".join(ch * 2 for ch in color)
    if len(color) != 6 or any(ch not in "0123456789abcdefABCDEF" for ch in color):
        raise ValueError(f"Invalid hex color: {value}")
    return RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))


def _luminance(color: RGBColor) -> float:
    return (0.299 * color[0] + 0.587 * color[1] + 0.114 * color[2]) / 255


def _tint(color: RGBColor, toward_white: float) -> RGBColor:
    ratio = min(max(toward_white, 0), 1)

    def mix(channel: int) -> int:
        return int(channel + (255 - channel) * ratio)

    return RGBColor(mix(color[0]), mix(color[1]), mix(color[2]))


def _on_color(color: RGBColor) -> RGBColor:
    return WHITE if _luminance(color) < 0.62 else INK


def _add_text_box(slide, name: str, left, top, width, height):
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.name = name
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    shape.line.fill.background()
    return shape


def _add_rect(slide, name: str, left, top, width, height, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _set_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _paint(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _footer_shapes(slide) -> None:
    _add_text_box(slide, "footer_text", Inches(0.62), Inches(7.05), Inches(8.4), Inches(0.28))
    number = _add_text_box(slide, "slide_number", Inches(10.15), Inches(7.05), Inches(2.55), Inches(0.28))
    number.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def _content_chrome(slide) -> None:
    _set_background(slide, PAPER)
    _add_rect(slide, "accent_bar", 0, 0, SLIDE_W, Inches(0.08), RGBColor(0x17, 0x32, 0x4D))
    _add_text_box(slide, "tag_text", Inches(0.62), Inches(0.32), Inches(11.8), Inches(0.28))
    _add_text_box(slide, "title_text", Inches(0.62), Inches(0.62), Inches(12.05), Inches(0.95))
    _add_rect(slide, "divider", Inches(0.62), Inches(1.66), Inches(1.55), Inches(0.045), RGBColor(0xC4, 0x6B, 0x3A))
    _footer_shapes(slide)


def _build_title_slide(slide) -> None:
    _set_background(slide, PAPER)
    _add_rect(slide, "accent_bar", 0, 0, Inches(7.2), SLIDE_H, RGBColor(0x17, 0x32, 0x4D))
    _add_rect(slide, "divider", Inches(7.2), 0, Inches(0.07), SLIDE_H, RGBColor(0xC4, 0x6B, 0x3A))
    _add_text_box(slide, "tag_text", Inches(0.55), Inches(1.7), Inches(6.1), Inches(0.32))
    _add_text_box(slide, "title_text", Inches(0.55), Inches(2.15), Inches(6.2), Inches(2.35))
    _add_text_box(slide, "body_text", Inches(0.55), Inches(4.7), Inches(6.15), Inches(1.45))
    _add_text_box(slide, "quote_text", Inches(7.7), Inches(2.55), Inches(5.05), Inches(2.6))
    _add_text_box(slide, "footer_text", Inches(7.7), Inches(6.95), Inches(3.3), Inches(0.28))
    number = _add_text_box(slide, "slide_number", Inches(10.9), Inches(6.95), Inches(1.85), Inches(0.28))
    number.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def _build_summary_slide(slide) -> None:
    _content_chrome(slide)
    _add_text_box(slide, "body_text", Inches(0.62), Inches(1.95), Inches(7.35), Inches(4.75))
    _add_rect(slide, "quote_panel", Inches(8.2), Inches(1.95), Inches(4.5), Inches(4.75), WHITE)
    _add_text_box(slide, "quote_text", Inches(8.45), Inches(2.18), Inches(4.05), Inches(4.3))


def _build_pain_slide(slide) -> None:
    _content_chrome(slide)
    _add_text_box(slide, "body_text", Inches(0.62), Inches(1.9), Inches(12.05), Inches(1.7))
    _add_rect(slide, "quote_panel", Inches(0.62), Inches(3.85), Inches(12.05), Inches(2.9), WHITE)
    mark = _add_text_box(slide, "quote_mark", Inches(0.82), Inches(3.95), Inches(0.7), Inches(0.6))
    mark.text_frame.paragraphs[0].text = "“"
    _add_text_box(slide, "quote_text", Inches(1.45), Inches(4.15), Inches(10.85), Inches(2.35))


def _build_solution_slide(slide) -> None:
    _content_chrome(slide)
    _add_text_box(slide, "body_text", Inches(0.62), Inches(1.88), Inches(12.05), Inches(0.55))
    _add_text_box(slide, "quote_text", Inches(0.62), Inches(6.55), Inches(12.05), Inches(0.38))


def _build_standard_slide(slide) -> None:
    _content_chrome(slide)
    _add_text_box(slide, "body_text", Inches(0.62), Inches(1.95), Inches(12.05), Inches(4.15))
    _add_text_box(slide, "quote_text", Inches(0.62), Inches(6.2), Inches(12.05), Inches(0.7))


def create_template(path: Path = TEMPLATE_PATH) -> Path:
    """Write the 6-slide template with named shapes."""
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = SLIDE_W
    presentation.slide_height = SLIDE_H
    blank = presentation.slide_layouts[6]
    builders = (
        _build_title_slide,
        _build_summary_slide,
        _build_pain_slide,
        _build_solution_slide,
        _build_standard_slide,
        _build_standard_slide,
    )
    for builder in builders:
        slide = presentation.slides.add_slide(blank)
        builder(slide)
    presentation.save(path)
    return path


def _shape_by_name(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    raise KeyError(name)


def _maybe(slide, name: str):
    try:
        return _shape_by_name(slide, name)
    except KeyError:
        return None


def _write(
    shape,
    text: str,
    color: RGBColor,
    size: int,
    bold: bool = False,
    italic: bool = False,
    align=PP_ALIGN.LEFT,
    space_after: int = 8,
) -> None:
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    lines = (text or "").split("\n") or [""]
    first = frame.paragraphs[0]
    first.clear()
    _fill_paragraph(first, lines[0], color, size, bold, italic, align, space_after)
    for extra in list(frame.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)
    for line in lines[1:]:
        paragraph = frame.add_paragraph()
        _fill_paragraph(paragraph, line, color, size, bold, italic, align, space_after)


def _fill_paragraph(paragraph, text, color, size, bold, italic, align, space_after) -> None:
    paragraph.alignment = align
    paragraph.space_after = Pt(space_after)
    paragraph.line_spacing = 1.08
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _write_pairs(shape, pairs: list[tuple[str, str]], label_color: RGBColor, value_color: RGBColor) -> None:
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    first = frame.paragraphs[0]
    first.clear()
    for extra in list(frame.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)

    def add(text: str, color: RGBColor, size: int, bold: bool, space_after: int, *, first_paragraph: bool = False):
        paragraph = frame.paragraphs[0] if first_paragraph else frame.add_paragraph()
        if first_paragraph:
            paragraph.clear()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color

    started = False
    for label, value in pairs:
        add(label, label_color, 11, True, 2, first_paragraph=not started)
        started = True
        add(value or "Not discussed", value_color, 16, False, 14)

def _recolor_runs(shape, color: RGBColor, size: int | None = None) -> None:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = color
            run.font.name = FONT
            if size:
                run.font.size = Pt(size)
            run.font.bold = True


def _duplicate_slide(presentation: Presentation, index: int):
    source = presentation.slides[index]
    destination = presentation.slides.add_slide(source.slide_layout)
    for shape in list(destination.shapes):
        element = shape._element
        element.getparent().remove(element)
    for shape in source.shapes:
        destination.shapes._spTree.append(deepcopy(shape._element))
    return destination


def _delete_slide(presentation: Presentation, index: int) -> None:
    slide_id = presentation.slides._sldIdLst[index]
    presentation.part.drop_rel(slide_id.get(qn("r:id")))
    presentation.slides._sldIdLst.remove(slide_id)


def _reorder(presentation: Presentation, order: list[int]) -> None:
    xml_slides = presentation.slides._sldIdLst
    existing = list(xml_slides)
    for child in list(xml_slides):
        xml_slides.remove(child)
    for index in order:
        xml_slides.append(existing[index])


def _prepare_pain_slides(presentation: Presentation, count: int) -> None:
    if count <= 0:
        _delete_slide(presentation, 2)
        return
    if count == 1:
        return
    for _ in range(count - 1):
        _duplicate_slide(presentation, 2)
    total = len(presentation.slides)
    extras = list(range(6, total))
    _reorder(presentation, [0, 1, 2, *extras, 3, 4, 5])


def _company_label(analysis: dict, fallback: str | None) -> str:
    name = (analysis.get("company_name") or fallback or "").strip()
    return name or "The customer"


def _add_logo(slide, logo_path: Path) -> None:
    from PIL import Image

    with Image.open(logo_path) as image:
        width_px, height_px = image.size
    if width_px <= 0 or height_px <= 0:
        return
    aspect = width_px / height_px
    height = Inches(1.35)
    width = Emu(int(height * aspect))
    max_width = Inches(4.4)
    if width > max_width:
        width = max_width
        height = Emu(int(width / aspect))
    if height > Inches(1.5):
        height = Inches(1.5)
        width = Emu(int(height * aspect))
    left = SLIDE_W - width - Inches(0.48)
    slide.shapes.add_picture(str(logo_path), left, Inches(0.42), width=width, height=height)


def _apply_brand(slide, primary: RGBColor, secondary: RGBColor) -> None:
    accent = _maybe(slide, "accent_bar")
    divider = _maybe(slide, "divider")
    panel = _maybe(slide, "quote_panel")
    mark = _maybe(slide, "quote_mark")
    if accent is not None:
        _paint(accent, primary)
    if divider is not None:
        _paint(divider, secondary)
    if panel is not None:
        _paint(panel, _tint(primary, 0.93))
    if mark is not None:
        _recolor_runs(mark, secondary, size=42)


def _stamp(slide, salesperson: str, number: int, total: int, *, light: bool = False) -> None:
    footer = _maybe(slide, "footer_text")
    counter = _maybe(slide, "slide_number")
    color = SOFT if light else MUTED
    who = salesperson.strip() or "Account team"
    if footer is not None:
        _write(footer, f"{who}   ·   Demo to Deck", color, 11, space_after=0)
    if counter is not None:
        _write(counter, f"{number:02d}  /  {total:02d}", color, 11, align=PP_ALIGN.RIGHT, space_after=0)


def _set_cell(cell, text: str, fill: RGBColor, font_color: RGBColor, *, bold: bool = False, size: int = 14, italic: bool = False) -> None:
    cell.text = ""
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.14)
    cell.margin_right = Inches(0.12)
    cell.margin_top = Inches(0.08)
    cell.margin_bottom = Inches(0.08)
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = font_color


def _solution_table(slide, pains: list[dict], primary: RGBColor, secondary: RGBColor) -> None:
    row_count = max(len(pains), 1) + 1
    table_height = Inches(0.58 * row_count)
    shape = slide.shapes.add_table(
        row_count,
        3,
        Inches(0.62),
        Inches(2.55),
        Inches(12.05),
        table_height,
    )
    table = shape.table
    table.columns[0].width = Inches(0.85)
    table.columns[1].width = Inches(5.35)
    table.columns[2].width = Inches(5.85)
    header = _on_color(primary)
    solve_fill = _tint(secondary, 0.88)
    _set_cell(table.cell(0, 0), "", primary, header, bold=True, size=12)
    _set_cell(table.cell(0, 1), "What they raised", primary, header, bold=True, size=13)
    _set_cell(table.cell(0, 2), "How we solve it", primary, header, bold=True, size=13)
    if not pains:
        _set_cell(table.cell(1, 0), "—", WHITE, MUTED, size=14)
        _set_cell(table.cell(1, 1), "No customer pain points were extracted from this recording.", WHITE, INK, size=14)
        _set_cell(table.cell(1, 2), "", solve_fill, MUTED, size=14)
        return
    for index, point in enumerate(pains, start=1):
        _set_cell(table.cell(index, 0), f"{index:02d}", WHITE, primary, bold=True, size=14)
        _set_cell(table.cell(index, 1), point.get("title") or "Pain point", WHITE, INK, size=15)
        _set_cell(table.cell(index, 2), "", solve_fill, MUTED, size=14)


def build_deck(
    analysis: dict,
    enrichment: dict,
    output_path: Path,
    *,
    salesperson_name: str,
    company_name: str | None,
    brand_color_primary: str,
    brand_color_secondary: str,
    logo_path: Path | None,
    template_path: Path = TEMPLATE_PATH,
) -> None:
    if not template_path.exists():
        create_template(template_path)
    primary = _rgb(brand_color_primary)
    secondary = _rgb(brand_color_secondary)
    on_primary = _on_color(primary)
    company = _company_label(analysis, company_name)
    pains = list(analysis.get("pain_points") or [])[:4]
    news = (enrichment or {}).get("recent_news_snippet")
    salesperson = salesperson_name or "the account team"

    presentation = Presentation(str(template_path))
    if len(presentation.slides) != 6:
        raise RuntimeError("Deck template does not contain the 6 base slides")
    _prepare_pain_slides(presentation, len(pains))

    title_slide = presentation.slides[0]
    summary_slide = presentation.slides[1]
    if pains:
        pain_slides = [presentation.slides[2 + index] for index in range(len(pains))]
        solution_slide = presentation.slides[2 + len(pains)]
        value_slide = presentation.slides[3 + len(pains)]
        next_slide = presentation.slides[4 + len(pains)]
    else:
        pain_slides = []
        solution_slide = presentation.slides[2]
        value_slide = presentation.slides[3]
        next_slide = presentation.slides[4]

    for slide in presentation.slides:
        _apply_brand(slide, primary, secondary)

    _write(_shape_by_name(title_slide, "tag_text"), "OPPORTUNITY BRIEF", secondary, 13, bold=True, space_after=0)
    _write(_shape_by_name(title_slide, "title_text"), company, on_primary, 40, bold=True, space_after=0)
    panel_body = SOFT if on_primary == WHITE else INK
    _write(
        _shape_by_name(title_slide, "body_text"),
        f"Prepared by {salesperson}\nA leave-behind from the call. The solve column is left blank on purpose.",
        panel_body,
        16,
        space_after=10,
    )
    if news:
        news_copy = f"Recently\n{news}"
    else:
        news_copy = "From the call\nPain points below use the customer's words. Nothing on the solve column is filled in for you."
    _write(_shape_by_name(title_slide, "quote_text"), news_copy, INK, 16, space_after=8)

    signals = analysis.get("buying_signals") or {}
    budget = signals.get("budget_detail") if signals.get("budget_mentioned") else "Not mentioned"
    timeline = signals.get("timeline_mentioned") or "Not mentioned"
    decision = "On the call" if signals.get("decision_maker_present") else "Not confirmed"
    _write(_shape_by_name(summary_slide, "tag_text"), "CALL SUMMARY", secondary, 12, bold=True, space_after=0)
    _write(_shape_by_name(summary_slide, "title_text"), "What we heard", primary, 32, bold=True, space_after=0)
    _write(
        _shape_by_name(summary_slide, "body_text"),
        analysis.get("call_summary") or "",
        INK,
        18,
        space_after=12,
    )
    _write_pairs(
        _shape_by_name(summary_slide, "quote_text"),
        [
            ("INDUSTRY", analysis.get("industry_guess") or "Not discussed"),
            ("BUDGET", budget or "Not mentioned"),
            ("TIMELINE", timeline),
            ("DECISION MAKER", decision),
        ],
        secondary,
        INK,
    )

    for slide, point in zip(pain_slides, pains):
        category = str(point.get("category", "other")).replace("_", " ").upper()
        severity = str(point.get("severity", "medium")).upper()
        _write(_shape_by_name(slide, "tag_text"), f"{severity}   ·   {category}", secondary, 12, bold=True, space_after=0)
        _write(_shape_by_name(slide, "title_text"), point.get("title") or "Pain point", primary, 30, bold=True, space_after=0)
        _write(_shape_by_name(slide, "body_text"), point.get("description") or "", INK, 18, space_after=8)
        _write(
            _shape_by_name(slide, "quote_text"),
            point.get("quote", "").strip(),
            INK,
            18,
            italic=True,
            space_after=0,
        )

    _write(_shape_by_name(solution_slide, "tag_text"), "SALESPERSON TO COMPLETE", secondary, 12, bold=True, space_after=0)
    _write(_shape_by_name(solution_slide, "title_text"), "Solution mapping", primary, 32, bold=True, space_after=0)
    _write(
        _shape_by_name(solution_slide, "body_text"),
        "Each problem from the call sits next to a blank. Fill in how you solve it. Do not add a product claim that was not yours to make.",
        MUTED,
        15,
        space_after=0,
    )
    _write(
        _shape_by_name(solution_slide, "quote_text"),
        "The right column stays empty until you write it.",
        MUTED,
        13,
        italic=True,
        space_after=0,
    )
    _solution_table(solution_slide, pains, primary, secondary)

    angle = analysis.get("recommended_solution_angle") or "Confirm the angle with the customer before proposing anything."
    high_count = sum(1 for point in pains if point.get("severity") == "high")
    value_body = "\n".join(
        [
            "Suggested conversation angle",
            angle,
            "",
            f"{len(pains)} pain points in this deck, {high_count} of them high severity.",
            "",
            "Customer impact to confirm",
            "________________________________",
            "",
            "ROI hypothesis, for you to complete",
            "________________________________",
            "",
            "Proof you will attach",
            "________________________________",
        ]
    )
    _write(_shape_by_name(value_slide, "tag_text"), "WORKING HYPOTHESIS", secondary, 12, bold=True, space_after=0)
    _write(_shape_by_name(value_slide, "title_text"), "Value and ROI", primary, 32, bold=True, space_after=0)
    _write(_shape_by_name(value_slide, "body_text"), value_body, INK, 16, space_after=6)
    _write(
        _shape_by_name(value_slide, "quote_text"),
        "Numbers on this slide are blank so the salesperson can complete them after the call.",
        MUTED,
        13,
        italic=True,
        space_after=0,
    )

    steps = [step for step in (analysis.get("next_steps") or []) if str(step).strip()]
    step_lines = [f"{index:02d}     {step}" for index, step in enumerate(steps, start=1)] or [
        "01     Agree the next meeting and the owner for each open question."
    ]
    objection_lines = []
    for item in analysis.get("objections") or []:
        handled = "Addressed on the call" if item.get("handled_on_call") else "Still open"
        objection_lines.append(f"{item.get('objection')}  —  {handled}")
    _write(_shape_by_name(next_slide, "tag_text"), "AGREED FOLLOW-UP", secondary, 12, bold=True, space_after=0)
    _write(_shape_by_name(next_slide, "title_text"), "Next steps", primary, 32, bold=True, space_after=0)
    _write(_shape_by_name(next_slide, "body_text"), "\n".join(step_lines), INK, 18, space_after=14)
    objection_copy = "\n".join(objection_lines) if objection_lines else "No objections were captured."
    _write(_shape_by_name(next_slide, "quote_text"), f"Objections\n{objection_copy}", MUTED, 14, space_after=4)

    if logo_path and Path(logo_path).exists():
        try:
            _add_logo(title_slide, Path(logo_path))
        except Exception:
            logger.warning("Logo could not be placed on the title slide", exc_info=True)

    total = len(presentation.slides)
    for index, slide in enumerate(presentation.slides, start=1):
        _stamp(slide, salesperson, index, total)

    presentation.core_properties.title = f"{company} opportunity brief"
    presentation.core_properties.author = salesperson_name or ""
    presentation.core_properties.subject = "Demo to Deck sales leave-behind"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output_path))
